"""Approximate slide renderer for QA without PowerPoint/LibreOffice.

Draws backgrounds, filled/outlined shapes, pictures and wrapped text (text inheritance from layout/master placeholders),
and reports text that overflows its box. Text is measured with Segoe UI widened by WIDEN to be conservative for
Montserrat. usage: python preview_deck.py deck.pptx out_dir [slide numbers...]
"""
import io
import os
import re
import sys

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

PX = 50.0            # pixels per cm
WIDEN = 1.10
FONTS = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts')
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
_fc = {}


def font(bold, size_px, mono=False):
    key = (bold, int(size_px), mono)
    if key not in _fc:
        name = ('consolab.ttf' if bold else 'consola.ttf') if mono else ('segoeuib.ttf' if bold else 'segoeui.ttf')
        _fc[key] = ImageFont.truetype(os.path.join(FONTS, name), max(4, int(size_px)))
    return _fc[key]


def cm(emu):
    return emu / 360000.0


def theme_colors(prs):
    part = [p for p in prs.part.package.iter_parts() if 'theme' in str(p.partname)][0]
    x = part.blob.decode('utf8')
    cols = dict(re.findall(r'<a:(dk1|lt1|dk2|lt2|accent\d)>.*?val="([0-9A-Fa-f]{6})"', x, re.S))
    m = {'tx1': cols.get('dk1', '000000'), 'bg1': cols.get('lt1', 'FFFFFF'), 'tx2': cols.get('dk2'), 'bg2': cols.get('lt2')}
    m.update({k: v for k, v in cols.items()})
    return m


def color_of(el, theme):
    """el: element containing a srgbClr/schemeClr child (e.g. a:solidFill)."""
    if el is None:
        return None
    c = el.find('a:srgbClr', NS)
    val = None
    if c is not None:
        val = c.get('val')
    else:
        c = el.find('a:schemeClr', NS)
        if c is not None:
            val = theme.get(c.get('val'), '000000')
    if val is None:
        return None
    rgb = [int(val[i:i + 2], 16) for i in (0, 2, 4)]
    lm = c.find('a:lumMod', NS)
    lo = c.find('a:lumOff', NS)
    if lm is not None:
        f = int(lm.get('val')) / 100000
        off = int(lo.get('val')) / 100000 if lo is not None else 0
        rgb = [min(255, int(v * f + 255 * off)) for v in rgb]
    alpha = c.find('a:alpha', NS)
    a = int(int(alpha.get('val')) / 100000 * 255) if alpha is not None else 255
    return tuple(rgb) + (a,)


def inherited(shape, slide, kind):
    """Default run properties (sz, bold, color element) for a placeholder from layout/master."""
    out = {}
    chain = []
    own = shape._element.find('.//a:lstStyle', NS)
    if own is not None:
        chain.append(own)
    if shape.is_placeholder:
        idx = shape.placeholder_format.idx
        ptype = shape.placeholder_format.type
        for src in (slide.slide_layout, slide.slide_layout.slide_master):
            for ph in src.placeholders:
                if ph.placeholder_format.idx == idx or (idx == 0 and ph.placeholder_format.type == ptype):
                    chain.append(ph._element)
                    break
        master = slide.slide_layout.slide_master._element
        style = master.find('.//p:titleStyle' if (idx == 0) else './/p:bodyStyle', NS)
        if style is not None:
            chain.append(style)
    for el in chain:
        d = el.find('.//a:lvl1pPr/a:defRPr', NS)
        if d is None:
            continue
        if 'sz' not in out and d.get('sz'):
            out['sz'] = int(d.get('sz')) / 100
        if 'b' not in out and d.get('b') is not None:
            out['b'] = d.get('b') == '1'
        if 'fill' not in out and d.find('a:solidFill', NS) is not None:
            out['fill'] = d.find('a:solidFill', NS)
    return out


def draw_text(img, d, shape, slide, theme, issues, sno):
    tf = shape.text_frame
    if not tf.text.strip():
        return
    body = shape._element.find('.//a:bodyPr', NS)
    x0, y0, w, h = cm(shape.left), cm(shape.top), cm(shape.width), cm(shape.height)
    li = cm(int(body.get('lIns', 91440))) if body is not None else 0.25
    ri = cm(int(body.get('rIns', 91440))) if body is not None else 0.25
    ti = cm(int(body.get('tIns', 45720))) if body is not None else 0.13
    bi = cm(int(body.get('bIns', 45720))) if body is not None else 0.13
    wrap = body is None or body.get('wrap') != 'none'
    anchor = body.get('anchor', 't') if body is not None else 't'
    inh = inherited(shape, slide, 'text')
    avail_w = (w - li - ri) * PX
    lines = []   # (segments [(text, font, color)], height_px)
    for p in tf.paragraphs:
        ppr = p._p.find('a:pPr', NS)
        lnspc = 1.0
        if ppr is not None and ppr.find('a:lnSpc/a:spcPct', NS) is not None:
            lnspc = int(ppr.find('a:lnSpc/a:spcPct', NS).get('val')) / 100000
        elif shape.is_placeholder:
            lnspc = 0.9
        spc_after = 0.0
        if ppr is not None and ppr.find('a:spcAft/a:spcPts', NS) is not None:
            spc_after = int(ppr.find('a:spcAft/a:spcPts', NS).get('val')) / 100
        words = []
        max_sz = 0
        for r in p.runs:
            rpr = r._r.find('a:rPr', NS)
            sz = int(rpr.get('sz')) / 100 if rpr is not None and rpr.get('sz') else inh.get('sz', 18)
            b = (rpr.get('b') == '1') if rpr is not None and rpr.get('b') is not None else inh.get('b', False)
            fill = rpr.find('a:solidFill', NS) if rpr is not None else None
            col = color_of(fill, theme) or color_of(inh.get('fill'), theme) or \
                ((242, 242, 242, 255) if shape.is_placeholder else (0, 0, 0, 255))
            latin = rpr.find('a:latin', NS) if rpr is not None else None
            mono = latin is not None and latin.get('typeface') in ('Consolas', 'Courier New')
            size_px = sz / 72 * 2.54 * PX
            f = font(b, size_px * (1 if mono else WIDEN), mono)
            fd = font(b, size_px, mono)
            max_sz = max(max_sz, size_px)
            for tok in re.split(r'(\s+)', r.text):
                if tok:
                    words.append((tok, f, fd, col))
        if not words:
            lines.append(([], (max_sz or 12) * 1.2 * lnspc, 0))
            continue
        cur, cur_w = [], 0.0
        for tok, f, fd, col in words:
            tw = f.getlength(tok)
            if wrap and cur and cur_w + tw > avail_w and not tok.isspace():
                lines.append((cur, max_sz * 1.2 * lnspc, 0))
                cur, cur_w = [], 0.0
            if not cur and tok.isspace():
                continue
            cur.append((tok, fd, col, tw)); cur_w += tw
            if not wrap and cur_w > avail_w:
                issues.append(f'slide {sno}: no-wrap text too wide: {tf.text[:40]!r}')
        lines.append((cur, max_sz * 1.2 * lnspc, spc_after / 72 * 2.54 * PX))
    total = sum(hh + sa for _, hh, sa in lines)
    box_h = (h - ti - bi) * PX
    if total > box_h + 0.15 * PX:
        issues.append(f'slide {sno}: OVERFLOW {total / PX:.2f} cm > {box_h / PX:.2f} cm: {tf.text[:60]!r}')
    y = (y0 + ti) * PX
    if anchor == 'ctr':
        y += (box_h - total) / 2
    elif anchor == 'b':
        y += box_h - total
    for segs, hh, sa in lines:
        xx = (x0 + li) * PX
        line_w = sum(s[3] for s in segs) / WIDEN
        algn = None
        if segs:
            xx_off = 0
        for tok, fd, col, tw in segs:
            d.text((xx, y + hh * 0.08), tok, font=fd, fill=col)
            xx += fd.getlength(tok)
        y += hh + sa


def render(prs, sno, slide, theme, issues):
    W, H = cm(prs.slide_width) * PX, cm(prs.slide_height) * PX
    img = Image.new('RGBA', (int(W), int(H)), (90, 20, 120, 255))
    # background image: layout first, then master
    for src in (slide.slide_layout, slide.slide_layout.slide_master):
        bg = src._element.find('.//p:bg//a:blip', NS)
        if bg is not None:
            blob = src.part.related_part(bg.get('{%s}embed' % NS['r'])).blob
            img.paste(Image.open(io.BytesIO(blob)).convert('RGBA').resize(img.size))
            break
    d = ImageDraw.Draw(img, 'RGBA')

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes)
                continue
            x0, y0 = cm(sh.left or 0) * PX, cm(sh.top or 0) * PX
            x1, y1 = x0 + cm(sh.width or 0) * PX, y0 + cm(sh.height or 0) * PX
            el = sh._element
            sppr = el.find('p:spPr', NS)
            if hasattr(sh, 'image'):
                try:
                    im = Image.open(io.BytesIO(sh.image.blob)).convert('RGBA').resize((max(1, int(x1 - x0)), max(1, int(y1 - y0))))
                    img.paste(im, (int(x0), int(y0)), im)
                except Exception:
                    pass
            if sppr is not None and sh.shape_type in (1, 9, 14, 17) or el.tag.endswith('cxnSp'):
                geom = sppr.find('a:prstGeom', NS) if sppr is not None else None
                prst = geom.get('prst') if geom is not None else 'rect'
                fill = color_of(sppr.find('a:solidFill', NS), theme) if sppr is not None else None
                if fill is None and sppr is not None and sppr.find('a:noFill', NS) is None and el.find('p:style/a:fillRef', NS) is not None \
                        and sh.shape_type == 1:
                    fill = color_of(el.find('p:style/a:fillRef', NS), theme)
                ln = sppr.find('a:ln', NS) if sppr is not None else None
                line = color_of(ln.find('a:solidFill', NS), theme) if ln is not None else None
                if ln is None and el.find('p:style/a:lnRef', NS) is not None and sh.shape_type == 1:
                    line = color_of(el.find('p:style/a:lnRef', NS), theme)
                if ln is not None and ln.find('a:noFill', NS) is not None:
                    line = None
                if prst == 'line' or el.tag.endswith('cxnSp') or sh.shape_type == 9:
                    if line:
                        d.line([(x0, y0), (x1, y1)], fill=line, width=2)
                elif fill or line:
                    r = 0
                    if prst == 'roundRect':
                        gd = geom.find('a:avLst/a:gd', NS)
                        adj = int(gd.get('fmla').split()[-1]) / 100000 if gd is not None else 0.1667
                        r = adj * min(x1 - x0, y1 - y0)
                    d.rounded_rectangle([x0, y0, x1, y1], radius=r, fill=fill, outline=line, width=2 if line else 0)
            if sh.has_text_frame:
                draw_text(img, d, sh, slide, theme, issues, sno)
            if x1 > W + 2 or y1 > H + 2 or x0 < -2:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    issues.append(f'slide {sno}: text shape outside slide: {sh.text_frame.text[:40]!r}')
    walk(slide.shapes)
    return img.convert('RGB')


def main():
    path, out = sys.argv[1], sys.argv[2]
    only = {int(v) for v in sys.argv[3:]}
    os.makedirs(out, exist_ok=True)
    prs = Presentation(path)
    theme = theme_colors(prs)
    issues = []
    for i, slide in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        render(prs, i, slide, theme, issues).save(os.path.join(out, f'slide_{i:02d}.png'))
    print('\n'.join(issues) if issues else 'no overflow detected')


if __name__ == '__main__':
    main()
